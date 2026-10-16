"""ทะเบียนรถ (สมาชิก/ผู้มาเยือน/เฝ้าระวัง), ประวัติ, รถที่อยู่ในพื้นที่ตอนนี้, นำเข้า/ส่งออก Excel"""
from __future__ import annotations

import io
import time
from datetime import datetime
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy import func, or_, select

from .. import constants as C
from ..db.models import Event, Vehicle, Visit
from ..events.service import audit
from ..export import TZ, th_date, to_xlsx
from ..plates.parser import format_display, normalize_plate_text, normalize_query
from ..plates.provinces import match_province
from .auth import CurrentUser, current_user, get_ctx, require_admin
from .queries import vehicle_dict, vehicle_stats, visit_dict, visits_in_range, today_range

router = APIRouter(prefix="/api", tags=["vehicles"])
XLSX_COLS = ["ทะเบียน", "จังหวัด", "ประเภทรถ", "สถานะ", "เจ้าของ", "บ้านเลขที่", "โทรศัพท์", "หมายเหตุ", "พบครั้งแรก",
             "พบล่าสุด"]


@router.get("/vehicles")
def list_vehicles(q: str | None = None, status: str | None = None, page: int = Query(1, ge=1),
                  size: int = Query(50, ge=1, le=500), ctx=Depends(get_ctx),
                  _: CurrentUser = Depends(current_user)):  # type: ignore[no-untyped-def]
    with ctx.db.session() as s:
        cond = []
        if q:
            nq = normalize_query(q)
            cond.append(or_(Vehicle.plate_norm.contains(nq) if nq else False, Vehicle.owner_name.contains(q),
                            Vehicle.house_no.contains(q), Vehicle.province.contains(q)))
        if status:
            cond.append(Vehicle.status.in_(status.split(",")))
        base = select(Vehicle).where(*cond)
        total = int(s.scalar(select(func.count()).select_from(base.subquery())) or 0)
        rows = s.scalars(base.order_by(Vehicle.last_seen.desc().nulls_last()).offset((page - 1) * size).limit(size))
        counts = dict(s.execute(select(Vehicle.status, func.count(Vehicle.id)).group_by(Vehicle.status)).all())
        return {"items": [vehicle_dict(v) for v in rows], "total": total, "counts": counts}


class VehicleIn(BaseModel):
    plate: str | None = None
    province: str | None = None
    vehicle_type: str | None = None
    status: str | None = None
    owner_name: str | None = None
    house_no: str | None = None
    phone: str | None = None
    note: str | None = None


def _apply(v: Vehicle, body: VehicleIn) -> list[str]:
    changes = []
    data = body.model_dump(exclude_unset=True)
    if "vehicle_type" in data and data["vehicle_type"] is not None:
        if data["vehicle_type"] and data["vehicle_type"] not in C.VEHICLE_TYPES:
            raise HTTPException(400, "ประเภทรถไม่ถูกต้อง")
        v.type_locked = bool(data["vehicle_type"])
    if "status" in data and data["status"] not in C.VEHICLE_STATUSES:
        raise HTTPException(400, "สถานะไม่ถูกต้อง")
    for k in ("vehicle_type", "status", "owner_name", "house_no", "phone", "note"):
        if k in data and data[k] is not None and getattr(v, k) != data[k]:
            changes.append(f"{k}: {getattr(v, k)} -> {data[k]}")
            setattr(v, k, data[k])
    return changes


@router.post("/vehicles")
def create_vehicle(body: VehicleIn, ctx=Depends(get_ctx), user: CurrentUser = Depends(require_admin)):  # type: ignore[no-untyped-def]
    norm, _, _ = normalize_plate_text(body.plate or "")
    if not norm:
        raise HTTPException(400, "กรุณากรอกทะเบียนรถ")
    pv = (body.province or "").strip()
    prov = (match_province(pv, 70)[0] or pv) if pv else ""
    with ctx.db.session() as s:
        if s.scalar(select(Vehicle).where(Vehicle.plate_norm == norm, Vehicle.province == prov)):
            raise HTTPException(409, "มีรถทะเบียนนี้อยู่แล้ว")
        v = Vehicle(plate_norm=norm, province=prov, status=body.status or C.VS_MEMBER, vehicle_type="")
        _apply(v, body)
        s.add(v)
        s.flush()
        # ผูกเหตุการณ์เก่าที่ยังไม่มีรถ
        for ev in s.scalars(select(Event).where(Event.plate_norm == norm, Event.vehicle_id.is_(None),
                                                or_(Event.province == prov, Event.province == ""))):
            ev.vehicle_id = v.id
            v.first_seen = min(v.first_seen or ev.ts, ev.ts)
            v.last_seen = max(v.last_seen or ev.ts, ev.ts)
        ctx.events.rebuild_visits(s, v.id)
        audit(s, user.username, "เพิ่มรถ", f"{format_display(norm)} {prov} ({v.status})")
        return vehicle_dict(v)


@router.post("/vehicles/from-event/{event_id}")
def from_event(event_id: int, body: VehicleIn, ctx=Depends(get_ctx), user: CurrentUser = Depends(require_admin)):  # type: ignore[no-untyped-def]
    """เพิ่ม/ตั้งสถานะรถจากเหตุการณ์ในคลิกเดียว"""
    with ctx.db.session() as s:
        ev = s.get(Event, event_id)
        if ev is None:
            raise HTTPException(404, "ไม่พบเหตุการณ์")
        if not ev.plate_norm:
            raise HTTPException(400, "เหตุการณ์นี้อ่านป้ายไม่ได้ กรุณาแก้ไขทะเบียนก่อน")
        v = s.get(Vehicle, ev.vehicle_id) if ev.vehicle_id else None
        if v is None:
            v = ctx.events.find_or_create_vehicle(s, ev.plate_norm, ev.province, ev.vehicle_type, ev.ts)
            ev.vehicle_id = v.id
        ch = _apply(v, body)
        audit(s, user.username, "แก้ไขข้อมูลรถ", f"{format_display(v.plate_norm)} {v.province}: {'; '.join(ch)}")
        return vehicle_dict(v)


@router.get("/vehicles/{vehicle_id}")
def get_vehicle(vehicle_id: int, ctx=Depends(get_ctx), _: CurrentUser = Depends(current_user)):  # type: ignore[no-untyped-def]
    with ctx.db.session() as s:
        v = s.get(Vehicle, vehicle_id)
        if v is None:
            raise HTTPException(404, "ไม่พบรถ")
        a, b = today_range()
        return {**vehicle_dict(v), "visits_today": visits_in_range(s, v.id, a, b),
                "stats_30d": vehicle_stats(s, v.id, 30)}


@router.patch("/vehicles/{vehicle_id}")
def patch_vehicle(vehicle_id: int, body: VehicleIn, ctx=Depends(get_ctx),
                  user: CurrentUser = Depends(require_admin)):  # type: ignore[no-untyped-def]
    with ctx.db.session() as s:
        v = s.get(Vehicle, vehicle_id)
        if v is None:
            raise HTTPException(404, "ไม่พบรถ")
        ch = _apply(v, body)
        if body.plate is not None or body.province is not None:
            norm = normalize_plate_text(body.plate)[0] if body.plate is not None else v.plate_norm
            pv = (body.province if body.province is not None else v.province).strip()
            prov = (match_province(pv, 70)[0] or pv) if pv else ""
            if (norm, prov) != (v.plate_norm, v.province):
                if s.scalar(select(Vehicle).where(Vehicle.plate_norm == norm, Vehicle.province == prov,
                                                  Vehicle.id != v.id)):
                    raise HTTPException(409, "มีรถทะเบียนนี้อยู่แล้ว")
                ch.append(f"ทะเบียน {format_display(v.plate_norm)} {v.province} -> {format_display(norm)} {prov}")
                v.plate_norm, v.province = norm, prov
        audit(s, user.username, "แก้ไขข้อมูลรถ", f"#{v.id}: {'; '.join(ch)}")
        out = vehicle_dict(v)
    ctx.hub.publish({"type": "vehicle_update", "vehicle": out})
    return out


@router.delete("/vehicles/{vehicle_id}")
def delete_vehicle(vehicle_id: int, ctx=Depends(get_ctx), user: CurrentUser = Depends(require_admin)):  # type: ignore[no-untyped-def]
    with ctx.db.session() as s:
        v = s.get(Vehicle, vehicle_id)
        if v is None:
            raise HTTPException(404, "ไม่พบรถ")
        for ev in s.scalars(select(Event).where(Event.vehicle_id == v.id)):
            ev.vehicle_id = None
        audit(s, user.username, "ลบรถ", f"{format_display(v.plate_norm)} {v.province}")
        s.delete(v)
    return {"ok": True}


@router.get("/vehicles/{vehicle_id}/visits")
def vehicle_visits(vehicle_id: int, limit: int = Query(100, le=1000), ctx=Depends(get_ctx),
                   _: CurrentUser = Depends(current_user)):  # type: ignore[no-untyped-def]
    with ctx.db.session() as s:
        rows = s.scalars(select(Visit).where(Visit.vehicle_id == vehicle_id)
                         .order_by(func.coalesce(Visit.entry_ts, Visit.exit_ts).desc()).limit(limit))
        return {"items": [visit_dict(v) for v in rows]}


@router.get("/visits/inside")
def inside_now(ctx=Depends(get_ctx), _: CurrentUser = Depends(current_user)):  # type: ignore[no-untyped-def]
    """รถที่อยู่ในพื้นที่ตอนนี้ (เข้าแล้วยังไม่ออก) และรายการ ไม่พบการออก"""
    with ctx.db.session() as s:
        rows = s.execute(select(Visit, Vehicle).join(Vehicle, Vehicle.id == Visit.vehicle_id)
                         .where(Visit.status.in_([C.VISIT_INSIDE, C.VISIT_NO_EXIT]), Visit.exit_ts.is_(None))
                         .order_by(Visit.entry_ts.desc()).limit(500)).all()
        inside, no_exit = [], []
        for visit, v in rows:
            item = {**visit_dict(visit), "vehicle": vehicle_dict(v)}
            (inside if visit.status == C.VISIT_INSIDE else no_exit).append(item)
        return {"inside": inside, "no_exit": no_exit[:100]}


@router.get("/vehicles-export.xlsx")
def export_vehicles(ctx=Depends(get_ctx), user: CurrentUser = Depends(current_user)):  # type: ignore[no-untyped-def]
    with ctx.db.session() as s:
        rows = [[format_display(v.plate_norm), v.province, v.vehicle_type, v.status, v.owner_name, v.house_no, v.phone,
                 v.note, th_date(v.first_seen, True) if v.first_seen else "", th_date(v.last_seen, True) if v.last_seen
                 else ""] for v in s.scalars(select(Vehicle).where(Vehicle.status != C.VS_UNKNOWN)
                                            .order_by(Vehicle.status, Vehicle.plate_norm))]
        audit(s, user.username, "ส่งออกข้อมูล", f"ทะเบียนรถ Excel {len(rows)} คัน")
    data = to_xlsx(XLSX_COLS, rows, "ทะเบียนรถ")
    name = f"lpr-vehicles-{datetime.now(TZ):%Y%m%d}.xlsx"
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.post("/vehicles-import")
async def import_vehicles(file: UploadFile = File(...), ctx=Depends(get_ctx),
                          user: CurrentUser = Depends(require_admin)):  # type: ignore[no-untyped-def]
    from openpyxl import load_workbook

    raw = await file.read()
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(400, "เปิดไฟล์ Excel ไม่ได้") from e
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(it, [])]
    idx = {name: header.index(name) for name in XLSX_COLS if name in header}
    if "ทะเบียน" not in idx:
        raise HTTPException(400, "ไม่พบคอลัมน์ 'ทะเบียน' ในแถวแรก (ดาวน์โหลดไฟล์ตัวอย่างจากปุ่มส่งออก)")
    added = updated = skipped = 0
    errors: list[str] = []

    def col(row: tuple[Any, ...], name: str) -> str:
        i = idx.get(name)
        return str(row[i]).strip() if i is not None and i < len(row) and row[i] is not None else ""

    with ctx.db.session() as s:
        for n, row in enumerate(it, start=2):
            norm, _, _ = normalize_plate_text(col(row, "ทะเบียน"))
            if not norm:
                skipped += 1
                continue
            pv = col(row, "จังหวัด")
            prov = (match_province(pv, 70)[0] or pv) if pv else ""
            status = col(row, "สถานะ") or C.VS_MEMBER
            if status not in C.VEHICLE_STATUSES:
                errors.append(f"แถว {n}: สถานะ '{status}' ไม่ถูกต้อง")
                continue
            vt = col(row, "ประเภทรถ")
            v = s.scalar(select(Vehicle).where(Vehicle.plate_norm == norm, Vehicle.province == prov))
            if v is None:
                v = Vehicle(plate_norm=norm, province=prov, status=status)
                s.add(v)
                added += 1
            else:
                updated += 1
            v.status = status
            if vt in C.VEHICLE_TYPES:
                v.vehicle_type, v.type_locked = vt, True
            for k, name in (("owner_name", "เจ้าของ"), ("house_no", "บ้านเลขที่"), ("phone", "โทรศัพท์"),
                            ("note", "หมายเหตุ")):
                val = col(row, name)
                if val:
                    setattr(v, k, val)
        audit(s, user.username, "นำเข้าทะเบียนรถ", f"เพิ่ม {added} แก้ไข {updated} ข้าม {skipped}")
    return {"added": added, "updated": updated, "skipped": skipped, "errors": errors[:50], "ts": time.time()}
