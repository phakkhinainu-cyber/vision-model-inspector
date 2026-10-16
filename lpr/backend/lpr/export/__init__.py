"""ส่งออกข้อมูล: CSV (UTF-8 BOM เปิดใน Excel ได้), Excel และรายงาน PDF (ฝังฟอนต์ Sarabun)"""
from __future__ import annotations

import csv
import io
from datetime import datetime
from typing import Any, Iterable
from zoneinfo import ZoneInfo

from ..config import resources_dir

TZ = ZoneInfo("Asia/Bangkok")
TH_MONTHS = ["", "ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]
COLOR_TH = {"white": "ขาว", "yellow": "เหลือง", "red": "แดง", "green": "เขียว", "blue": "น้ำเงิน", "black": "ดำ",
            "orange": "ส้ม", "unknown": "-"}

EVENT_COLUMNS = ["วันที่", "เวลา", "ทิศทาง", "ทะเบียน", "จังหวัด", "ประเภทรถ", "สีป้าย", "สถานะรถ", "เจ้าของ",
                 "บ้านเลขที่", "กล้อง", "ความแม่นยำ (%)", "การตรวจสอบ", "ที่มา", "แก้ไขโดย"]


def th_date(ts: float, short: bool = False) -> str:
    d = datetime.fromtimestamp(ts, TZ)
    if short:
        return f"{d.day:02d}/{d.month:02d}/{d.year + 543}"
    return f"{d.day} {TH_MONTHS[d.month]} {d.year + 543}"


def th_time(ts: float) -> str:
    return datetime.fromtimestamp(ts, TZ).strftime("%H:%M:%S")


def event_row(e: dict[str, Any]) -> list[Any]:
    src = {"auto": "อัตโนมัติ", "capture": "จับภาพเดี๋ยวนี้", "manual": "บันทึกด้วยมือ"}.get(e.get("source", ""), "")
    color = f"{COLOR_TH.get(e.get('plate_bg') or 'unknown', '-')}/{COLOR_TH.get(e.get('plate_fg') or 'unknown', '-')}"
    return [th_date(e["ts"], short=True), th_time(e["ts"]), e["direction"], e.get("plate_display") or "(อ่านไม่ได้)",
            e.get("province") or "", e.get("vehicle_type") or "", color, e.get("vehicle_status") or "",
            e.get("owner_name") or "", e.get("house_no") or "", e.get("camera_name") or "",
            round(100 * float(e.get("ocr_confidence") or 0)), e.get("review_status") or "", src,
            e.get("edited_by") or ""]


def to_csv(columns: list[str], rows: Iterable[list[Any]]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(columns)
    for r in rows:
        w.writerow(r)
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(columns: list[str], rows: Iterable[list[Any]], title: str = "ข้อมูล") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    ws.append(columns)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4338CA")
        c.alignment = Alignment(horizontal="center")
    widths = [len(c) + 2 for c in columns]
    for r in rows:
        ws.append(r)
        for i, v in enumerate(r):
            widths[i] = max(widths[i], min(40, len(str(v)) + 2))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(title: str, subtitle: str, summary: list[tuple[str, str]], events: list[dict[str, Any]],
           max_rows: int = 3000) -> bytes:
    from fpdf import FPDF

    fonts = resources_dir() / "fonts"
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.add_font("Sarabun", "", str(fonts / "Sarabun-Regular.ttf"))
    pdf.add_font("Sarabun", "B", str(fonts / "Sarabun-Bold.ttf"))
    try:
        pdf.set_text_shaping(True)  # จัดวางสระ/วรรณยุกต์ไทยให้ถูกต้อง (ใช้ uharfbuzz)
    except Exception:  # noqa: BLE001
        pass
    pdf.set_auto_page_break(True, margin=12)
    pdf.set_title(title)
    pdf.set_creator("LPR บ้าน")
    pdf.add_page()
    pdf.set_font("Sarabun", "B", 18)
    pdf.cell(0, 10, title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Sarabun", "", 11)
    pdf.set_text_color(86, 95, 120)
    pdf.cell(0, 7, subtitle, new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(20, 26, 46)
    pdf.ln(2)
    # สรุป
    pdf.set_font("Sarabun", "", 11)
    col_w = (pdf.w - pdf.l_margin - pdf.r_margin) / max(1, min(6, len(summary)))
    for i, (k, v) in enumerate(summary[:6]):
        x = pdf.l_margin + i * col_w
        pdf.set_xy(x, pdf.get_y())
        pdf.set_fill_color(238, 241, 247)
        pdf.cell(col_w - 2, 7, k, fill=True)
        pdf.set_xy(x, pdf.get_y() + 7)
        pdf.set_font("Sarabun", "B", 14)
        pdf.cell(col_w - 2, 9, v, fill=True)
        pdf.set_font("Sarabun", "", 11)
        pdf.set_xy(x, pdf.get_y() - 7)
    pdf.ln(20)
    cols = ["วันที่", "เวลา", "ทิศทาง", "ทะเบียน", "จังหวัด", "ประเภทรถ", "สถานะรถ", "เจ้าของ", "กล้อง", "แม่นยำ", "ตรวจสอบ"]
    widths = [22, 18, 14, 28, 34, 26, 22, 38, 30, 16, 25]
    scale = (pdf.w - pdf.l_margin - pdf.r_margin) / sum(widths)
    widths = [w * scale for w in widths]

    def header() -> None:
        pdf.set_font("Sarabun", "B", 10)
        pdf.set_fill_color(67, 56, 202)
        pdf.set_text_color(255, 255, 255)
        for c, w in zip(cols, widths):
            pdf.cell(w, 7, c, border=0, fill=True)
        pdf.ln(7)
        pdf.set_text_color(20, 26, 46)
        pdf.set_font("Sarabun", "", 9.5)

    header()
    for i, e in enumerate(events[:max_rows]):
        if pdf.get_y() > pdf.h - 18:
            pdf.add_page()
            header()
        r = event_row(e)
        vals = [r[0], r[1], r[2], r[3], r[4], r[5], r[7], r[8], r[10], f"{r[11]}%", r[12]]
        fill = i % 2 == 1
        pdf.set_fill_color(246, 247, 251)
        if e["direction"] == "เข้า":
            dir_color = (29, 78, 216)
        else:
            dir_color = (194, 65, 12)
        for j, (v, w) in enumerate(zip(vals, widths)):
            if j == 2:
                pdf.set_text_color(*dir_color)
            pdf.cell(w, 6.2, _fit(pdf, str(v), w), fill=fill)
            if j == 2:
                pdf.set_text_color(20, 26, 46)
        pdf.ln(6.2)
    if len(events) > max_rows:
        pdf.ln(3)
        pdf.cell(0, 6, f"แสดง {max_rows:,} จาก {len(events):,} รายการ (ส่งออก Excel เพื่อดูทั้งหมด)")
    return bytes(pdf.output())


def _fit(pdf, text: str, w: float) -> str:  # type: ignore[no-untyped-def]
    if pdf.get_string_width(text) <= w - 1.5:
        return text
    while text and pdf.get_string_width(text + "…") > w - 1.5:
        text = text[:-1]
    return text + "…"
